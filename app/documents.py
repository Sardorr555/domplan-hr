from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor

TITLES = {'manager':'Руководитель отдела','hr':'HR','ceo':'Генеральный директор','accountant':'Бухгалтерия','superadmin':'Суперадминистратор','payroll':'Расчёт зарплаты','calculated':'Ожидается запись выплаты','paid':'Выплата записана','rejected':'Отклонено'}

def build_workflow(case=None, history=None, attachments=None):
    d=Document(); s=d.sections[0]
    s.top_margin=s.bottom_margin=Inches(.65)
    normal=d.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    for name in ['Title','Heading 1','Heading 2']:
        d.styles[name].font.color.rgb=RGBColor.from_string('1C6650')
    s.header.paragraphs[0].text='DOMPLAN  /  PEOPLE OPERATIONS'
    d.add_heading('Приём сотрудника и первая зарплата',0)
    d.add_paragraph('Единая маршрутная карточка • '+('Заявка № '+str(case['id']) if case else 'Шаблон для заполнения'))
    d.add_paragraph('Руководитель → HR → Директор → Бухгалтерия → Расчёт → Запись выплаты')
    if case and case.get('autonomous_demo'):
        d.add_paragraph('УЧЕБНАЯ СИМУЛЯЦИЯ. Реальный приём и банковский перевод не выполнялись.')
        if case.get('ai_pending'):
            proposal=case['ai_pending'];d.add_heading('Ожидается подтверждение человека',1);d.add_paragraph(proposal['conclusion'])
            for check in proposal['checks']:d.add_paragraph(check)
    if case:d.add_paragraph('Текущий этап: '+TITLES.get(case['state'],case['state'])+' | Версия: '+str(case['version']))
    def table(rows):
        t=d.add_table(rows=0,cols=2);t.style='Light Shading Accent 1'
        for k,v in rows:
            cells=t.add_row().cells;cells[0].text=k;cells[1].text=str(v if v is not None else '________________________')
    c=case or {}
    d.add_heading('1. Данные сотрудника',1)
    table([('ФИО',c.get('name')),('Отдел',c.get('department')),('Должность',c.get('position')),('Дата начала работы',c.get('start_date')),('Месячный оклад / валюта',str(c['salary'])+' / указать валюту' if c else None)])
    d.add_heading('2. Руководитель отдела',1)
    table([('Результат собеседования и рекомендация',c.get('conclusion')),('Ответственный / дата',None),('Решение: утвердить / вернуть / отклонить',None)])
    d.add_heading('3. HR: проверка и оформление',1)
    for t in ['Документы личности проверены','Сведения кандидата сверены','Проект договора и документы приёма подготовлены','Недостающие документы / замечания указаны']:
        d.add_paragraph('☐ '+t)
    table([('Заключение HR',None),('Ответственный / дата / решение',None)])
    d.add_heading('4. Генеральный директор',1)
    table([('Заключения руководителя и HR рассмотрены',None),('Утверждённые должность, оклад, дата начала',None),('Решение и обоснование',None),('Ответственный / дата',None)])
    d.add_heading('5. Бухгалтерия: первая зарплата',1)
    p=c.get('payroll',{})
    table([('Период',p.get('period')),('Рабочих дней по графику',p.get('scheduled_days')),('Отработано дней',p.get('worked_days')),('Премия',p.get('bonus')),('Удержания, включая налоги',p.get('deduction')),('Начислено',p.get('gross')),('К выплате',p.get('net')),('Плановая дата выплаты',p.get('payment_due'))])
    d.add_paragraph('Начислено = оклад × отработанные дни / дни по графику + премия; к выплате = начислено − удержания. Все суммы и удержания проверяет бухгалтер.')
    payment=c.get('payment',{})
    table([('Дата фактической внешней выплаты',payment.get('paid_on')),('Номер платёжного документа',payment.get('reference')),('Ответственный / дата',None)])
    d.add_heading('6. Приложенные документы',1)
    if attachments:
        for a in attachments:d.add_paragraph(a['name']+' | '+a['at'])
    else:d.add_paragraph('Укажите названия и версии приложений: ________________________')
    d.add_heading('7. История и контроль',1)
    if history:
        for a in history:
            d.add_paragraph(a['at']+' | '+a['actor']+' | '+a['action'])
            d.add_paragraph(a['detail'])
    else:table([('Этап / ответственный / дата / решение / комментарий',None),('Возврат: причина и необходимые исправления',None)])
    d.add_paragraph('Суперадминистратор может выполнить действие любой роли. Его решения также фиксируются в истории.')
    d.add_paragraph('Word-файл является снимком или рабочим шаблоном. Источник текущего статуса и согласований — приложение. Изменение файла не меняет статус заявки. При необходимости загрузите исправленный файл новой версией и оформите решение в системе.')
    d.add_paragraph('Запись выплаты не выполняет денежный перевод. Карточка не заменяет трудовой договор и обязательное оформление.')
    s.footer.paragraphs[0].text='DOMPLAN HR • Внутренний документ • Содержит персональные данные'
    out=BytesIO();d.save(out);return out.getvalue()
