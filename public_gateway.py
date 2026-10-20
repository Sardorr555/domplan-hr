from run import configure
configure()
from pathlib import Path
from io import BytesIO
from contextlib import asynccontextmanager
from datetime import datetime, timezone, timedelta
from urllib.parse import urlsplit
import json,os
from fastapi import FastAPI, Request, BackgroundTasks, HTTPException
from fastapi.responses import FileResponse, JSONResponse
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from app.main import DB, Base, engine, word_response
from app.demo import DemoRun, output, fictional, execute

STATIC=Path(__file__).parent/'app'/'static'
@asynccontextmanager
async def lifespan(app):
    Base.metadata.create_all(engine)
    yield
public_app=FastAPI(docs_url=None,redoc_url=None,openapi_url=None,lifespan=lifespan)
@public_app.middleware('http')
async def security(request,call_next):
    response=await call_next(request)
    response.headers.update({'Cache-Control':'no-store','X-Content-Type-Options':'nosniff','X-Frame-Options':'DENY','Content-Security-Policy':"default-src 'self'; script-src 'self'; style-src 'self'; frame-ancestors 'none'"})
    return response
@public_app.get('/')
def home():return FileResponse(STATIC/'public-demo.html')
@public_app.get('/public-demo.js')
def script():return FileResponse(STATIC/'public-demo.js',media_type='text/javascript')
@public_app.get('/public-demo.css')
def style():return FileResponse(STATIC/'public-demo.css',media_type='text/css')
@public_app.get('/api/runs')
def runs():
    with DB() as db:return [output(r) for r in db.scalars(select(DemoRun).order_by(DemoRun.id.desc()).limit(5))]
@public_app.post('/api/runs')
async def start(request:Request,tasks:BackgroundTasks):
    origin=request.headers.get('origin','')
    if urlsplit(origin).netloc!=request.headers.get('host') or not request.headers.get('content-type','').startswith('application/json'):
        raise HTTPException(403,'Обновите страницу и запустите демонстрацию с сайта.')
    if not os.getenv('RAGFLOW_API_KEY'):raise HTTPException(503,'Подключение Swipies не настроено.')
    with DB() as db:
        existing=db.scalar(select(DemoRun).where(DemoRun.status=='running'))
        if existing:return JSONResponse(output(existing),status_code=200)
        since=(datetime.now(timezone.utc)-timedelta(hours=1)).isoformat()
        recent=list(db.scalars(select(DemoRun.id).where(DemoRun.created>=since)))
        if len(recent)>=6:raise HTTPException(429,'Лимит демонстраций на этот час достигнут. Посмотрите готовые результаты или попробуйте позже.')
        r=DemoRun(status='running',slot=1,stage='manager',data=json.dumps({'card':fictional(),'events':[]},ensure_ascii=False));db.add(r)
        try:db.commit()
        except IntegrityError:db.rollback();raise HTTPException(409,'Агент уже запущен. Обновите страницу.')
        tasks.add_task(execute,r.id);return output(r)
@public_app.get('/api/runs/{id}/document.docx')
def document(id:int):
    with DB() as db:return build_word(id,db)

def build_word(id,db):
    from docx import Document
    r=db.get(DemoRun,id)
    if not r:raise HTTPException(404,'Демонстрация не найдена')
    data=json.loads(r.data);d=Document();d.add_heading('DOMPLAN • АВТОНОМНОЕ ДЕМО ИИ',0)
    d.add_paragraph('Все данные вымышленные. Документ не создаёт реальных кадровых решений и платежей.')
    card=data['card'];d.add_heading(card['name'],1)
    for text in [card['position'],card['department'],'Условный оклад: '+card['salary'],*card['documents']]:d.add_paragraph(text)
    for a in data['events']:
        d.add_heading(a['actor'],1);d.add_paragraph(a['at'][:16].replace('T',' '))
        d.add_paragraph({'approve':'Передано дальше','return':'Возвращено на повторную проверку','reject':'Отклонено в симуляции','simulated_payment':'Имитация выплаты'}.get(a['decision'],'Проверено'))
        d.add_paragraph(a['conclusion'])
        for check in a['checks']:d.add_paragraph(check,style='List Bullet')
    if data.get('payroll'):d.add_heading('К выплате в симуляции: '+data['payroll']['net'],1)
    out=BytesIO();d.save(out);return word_response(out.getvalue(),'demo-agent-'+str(id)+'.docx')

if __name__=='__main__':
    from run import configure
    configure()
    import uvicorn
    uvicorn.run('public_gateway:public_app',host='127.0.0.1',port=8001)
